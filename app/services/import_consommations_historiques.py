"""Import de consommations/pesées historiques, transcrites depuis des
fiches manuscrites -- demandé le 28/07/2026.

Contexte : les consommations de sources étaient jusqu'ici tenues sur des
fiches papier, jamais saisies dans l'appli. Plutôt qu'un bouton d'import
dans l'interface, le flux retenu est volontairement simple :

1. Les fiches papier sont transcrites (par Claude, à partir de scans)
   dans un fichier Excel, au format décrit ci-dessous.
2. Cette transcription est relue et corrigée par l'utilisateur --
   c'est LE moment de vérification, pas l'import lui-même.
3. Le fichier relu est importé avec ce script : un import "idiot",
   volontairement peu sophistiqué, puisque le contrôle a déjà eu lieu à
   l'étape 2.

Format attendu (une ligne par événement, en-têtes en première ligne) :

    ID Source          -- obligatoire, doit exister déjà dans l'appli
    Date               -- obligatoire (JJ/MM/AAAA ou AAAA-MM-JJ)
    Masse avant (g)    -- optionnel (mode pesée -- liquide)
    Masse après (g)    -- optionnel (mode pesée -- vide = pesée de
                          contrôle, masse avant reprise telle quelle)
    Quantité utilisée  -- optionnel (mode simple -- gaz, ou saisie
                          directe sans pesée)
    Commentaire        -- optionnel, texte libre (transcription y
                          compris les mentions illisibles/incertaines)
    Utilisateur        -- optionnel, si identifiable sur la fiche

Au moins l'un de (Masse avant, Quantité utilisée) doit être renseigné
par ligne -- sinon, rien à enregistrer. L'ordre des lignes n'a AUCUNE
importance (ni entre elles, ni par rapport aux consommations déjà en
base) : quantite_restante_calculee() trie toujours par date au moment du
calcul, jamais par ordre d'insertion (vérifié dans units.py).
"""
from datetime import datetime, date as date_type
from typing import Optional

import openpyxl

from app.models.consumption import ConsumptionDB
from app.models.source import SourceDB
from app.repositories.user import UserRepository


def _to_float(val) -> Optional[float]:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(",", ".").strip())
    except ValueError:
        return None


def _to_date(val):
    if val is None or val == "":
        return None
    if isinstance(val, (datetime, date_type)):
        return val
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(val).strip(), fmt)
        except ValueError:
            continue
    return None


def importer_consommations_historiques(fichier_path: str, db, utilisateur_par_defaut: str = None) -> dict:
    """Lit le fichier Excel transcrit et crée les consommations
    historiques correspondantes. Renvoie un rapport (dict) : nombre
    créées, et le détail de chaque ligne ignorée avec sa raison -- la
    seule "vérification" faite ici, le contrôle de fond ayant déjà eu
    lieu lors de la relecture de la transcription (voir le docstring du
    module).

    Détection de doublons potentiels ajoutée le 31/07/2026, suite à un
    cas réel (le même document importé deux fois par erreur, dupliquant
    tous les prélèvements) : avant chaque création, vérifie si une
    consommation existe déjà pour la MÊME source à la MÊME date (déjà en
    base, ou déjà créée plus tôt dans ce même import) -- signalé dans le
    rapport, mais n'empêche jamais la création (ça reste un import
    "idiot" : deux prélèvements réels le même jour sur la même source
    sont possibles, ce n'est pas forcément une erreur). Rien de plus
    simple ensuite que de supprimer le doublon depuis la page
    Consommations si le signalement s'avère justifié.

    Rattachement utilisateur ajouté le 31/08/2026 : la colonne
    Utilisateur, quand elle est renseignée, est rattachée à un compte
    existant (recherche insensible à la casse) ou à un enregistrement
    historique nouvellement créé -- jamais laissée comme texte seul sans
    lien, même si la personne n'a pas de compte et n'en aura jamais (cas
    réel : quelqu'un qui a quitté le service depuis longtemps). Chaque
    enregistrement historique créé est listé dans le rapport, pour être
    relu et éventuellement fusionné avec un compte existant depuis la
    fiche utilisateur si le rapprochement automatique a raté une
    variante orthographique.
    """
    rapport = {"consommations_creees": 0, "lignes_ignorees": [], "doublons_potentiels": [], "utilisateurs_historiques_crees": []}

    wb = openpyxl.load_workbook(fichier_path, data_only=True)
    ws = wb.active
    entetes = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(entetes) if h}

    def val(row, nom):
        idx = col.get(nom)
        if idx is None or idx >= len(row):
            return None
        v = row[idx].value if hasattr(row[idx], "value") else row[idx]
        return v if v not in (None, "") else None

    # (source_id, date) déjà rencontrés -- en base AVANT cet import, et
    # au fil de celui-ci (deux lignes du même fichier peuvent aussi se
    # dupliquer entre elles, pas seulement par rapport à l'existant).
    dates_existantes = {
        (c.source_id, c.timestamp.date())
        for c in db.query(ConsumptionDB).all()
        if c.timestamp is not None
    }

    for num_ligne, row in enumerate(ws.iter_rows(min_row=2), start=2):
        source_id = val(row, "ID Source")
        if not source_id:
            continue  # ligne vide, silencieusement ignorée (pas une erreur)
        source_id = str(source_id).strip()

        source = db.query(SourceDB).filter(SourceDB.id == source_id).first()
        if not source:
            rapport["lignes_ignorees"].append(
                {"ligne": num_ligne, "source_id": source_id, "raison": f"source '{source_id}' introuvable"}
            )
            continue

        date_evenement = _to_date(val(row, "Date"))
        if date_evenement is None:
            rapport["lignes_ignorees"].append(
                {"ligne": num_ligne, "source_id": source_id, "raison": "date manquante ou illisible"}
            )
            continue

        masse_avant = _to_float(val(row, "Masse avant (g)"))
        masse_apres = _to_float(val(row, "Masse après (g)"))
        quantite_utilisee = _to_float(val(row, "Quantité utilisée"))

        if masse_avant is None and quantite_utilisee is None:
            rapport["lignes_ignorees"].append(
                {"ligne": num_ligne, "source_id": source_id,
                 "raison": "ni masse avant ni quantité utilisée renseignées"}
            )
            continue

        cle = (source_id, date_evenement.date())
        if cle in dates_existantes:
            rapport["doublons_potentiels"].append({
                "ligne": num_ligne, "source_id": source_id, "date": date_evenement.date(),
            })
        dates_existantes.add(cle)

        if masse_avant is not None:
            # Pesée de contrôle si "masse après" absente -- même
            # convention que la saisie normale (voir ConsumptionService).
            if masse_apres is None:
                masse_apres = masse_avant
            quantite_finale = masse_avant - masse_apres
        else:
            quantite_finale = quantite_utilisee

        commentaire = val(row, "Commentaire")
        utilisateur = val(row, "Utilisateur") or utilisateur_par_defaut
        utilisateur = str(utilisateur).strip() if utilisateur else None

        utilisateur_id = None
        if utilisateur:
            deja_connu = UserRepository(db).get_by_username_insensible_casse(utilisateur) is not None
            utilisateur_lie = UserRepository(db).get_ou_creer_historique(utilisateur)
            utilisateur_id = utilisateur_lie.id
            if not deja_connu and utilisateur_lie.username not in rapport["utilisateurs_historiques_crees"]:
                rapport["utilisateurs_historiques_crees"].append(utilisateur_lie.username)

        db.add(ConsumptionDB(
            source_id=source_id,
            timestamp=date_evenement,
            quantite_utilisee=quantite_finale,
            masse_avant=masse_avant,
            masse_apres=masse_apres,
            commentaire=str(commentaire) if commentaire else None,
            utilisateur=utilisateur,
            utilisateur_id=utilisateur_id,
        ))
        rapport["consommations_creees"] += 1

    db.commit()
    return rapport


def generer_modele_vide(output_path: str) -> None:
    """Génère un fichier Excel vide avec les bons en-têtes et une ligne
    d'exemple, pour cadrer la transcription."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consommations historiques"
    ws.append(["ID Source", "Date", "Masse avant (g)", "Masse après (g)", "Quantité utilisée", "Commentaire", "Utilisateur"])
    ws.append(["SCA-0001", "15/03/2019", "125.4", "123.1", "", "Prélèvement pour étalonnage banc ICARE", "J. Dupont"])
    ws.append(["SCA-0002", "22/06/2019", "", "", "0.5", "Fiche partiellement illisible, quantité estimée", ""])
    for col_lettre, largeur in zip("ABCDEFG", (14, 12, 16, 16, 16, 40, 14)):
        ws.column_dimensions[col_lettre].width = largeur
    wb.save(output_path)
