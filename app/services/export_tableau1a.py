"""Génère le Tableau 1a (Récapitulatif des stocks physiques de matières
nucléaires par matière et par lieu) à partir des mêmes masses que l'Annexe
1 — le tableau 1a est une réorganisation de l'Annexe 1, sommée par matière
et triée par lieu (indiqué le 10/07/2026).

Limites connues :
- Seuls les lieux dont le "site" vaut IRMA ou EPICEA (comparaison insensible
  à la casse) sont reconnus. Une source Matière Nucléaire ailleurs (ou sans
  emplacement actuel suivi) est comptée dans le Total mais signalée à part
  (voir la valeur de retour `avertissements`).
- Deutérium et Lithium-6 enrichi sont des isotopes STABLES (non
  radioactifs) : leur masse ne peut pas être déduite d'une activité (qui
  n'existe pas pour un isotope stable). Ces deux lignes du tableau restent
  donc toujours à zéro avec le système actuel (basé sur l'activité) — il
  faudrait un moyen de saisir directement leur masse pour les couvrir.
- La distinction naturel / enrichi / appauvri est déduite du % massique
  d'U-235 calculé, avec une tolérance de ±0,05 point autour de la
  composition naturelle exacte (voir matieres_nucleaires.categoriser_uranium)
  — une approximation numérique, pas une donnée déclarée par source.
"""
from pathlib import Path
from collections import defaultdict
import openpyxl

from app.services.matieres_nucleaires import (
    parser_radionuclide, masse_radionuclide_g, categoriser_uranium,
)
from app.services.units import activite_actuelle_bq

TEMPLATE_PATH = Path(__file__).parent.parent / "resources" / "tableau1a_template.xlsx"

# (ligne dans le gabarit, unité de conversion depuis les grammes calculés, fusionné ?)
LIGNES = {
    "plutonium": (7, 1, True),
    "enrichi_20_plus": (8, 1, False),
    "enrichi_10_20": (9, 1, False),
    "enrichi_moins_10": (10, 1, False),
    "uranium_233": (11, 1, True),
    "naturel": (12, 1000, True),   # en kilogramme
    "appauvri": (13, 1000, True),  # en kilogramme
    "thorium": (14, 1000, True),   # en kilogramme
    "deuterium": (15, 1, True),
    "tritium": (16, 1, True),
    "lithium6": (17, 1, True),
}
COLONNES_LOCALISATION = {"EPICEA": "C", "IRMA": "E", "TOTAL": "G"}


def generer_tableau1a(db, output_path: str) -> dict:
    """Renvoie un rapport {avertissements: [...], sources_incluses: N}.

    Sources archivées exclues du comptage (01/09/2026) -- voir
    export_annexe1.py pour le raisonnement complet, appliqué ici de la
    même façon."""
    from app.repositories.source import SourceRepository
    from app.models.source import is_archived

    sources = [s for s in SourceRepository(db).get_all() if s.matiere_nucleaire and not is_archived(s)]

    # totals[categorie]["EPICEA"|"IRMA"|"AUTRE"]["total"|"u235"] = masse en grammes
    totals = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    avertissements = []
    sources_incluses = 0

    for source in sources:
        if not source.radionuclides:
            continue

        site = (source.emplacement_actuel.site or "").strip().upper() if source.emplacement_actuel else None
        if site == "EPICEA":
            lieu = "EPICEA"
        elif site == "IRMA":
            lieu = "IRMA"
        else:
            lieu = "AUTRE"
            avertissements.append(
                f"{source.id} : emplacement actuel non reconnu comme IRMA ou EPICEA "
                f"({source.emplacement_actuel.nom if source.emplacement_actuel else 'aucun emplacement suivi'}) "
                "— comptée dans le Total uniquement."
            )

        masses_uranium = {}  # {nombre_masse: masse_g}, hors U-233
        source_incluse = False

        for rn in source.radionuclides:
            if not rn.periode:
                continue
            parsed = parser_radionuclide(rn.nom)
            if not parsed:
                continue
            symbole, nombre_masse = parsed

            activite_actuelle = activite_actuelle_bq(rn)
            masse_g = masse_radionuclide_g(db, rn, activite_actuelle)
            if masse_g is None:
                continue

            if symbole == "Pu":
                totals["plutonium"][lieu]["total"] += masse_g
                totals["plutonium"]["TOTAL"]["total"] += masse_g
                source_incluse = True
            elif symbole == "Th":
                totals["thorium"][lieu]["total"] += masse_g
                totals["thorium"]["TOTAL"]["total"] += masse_g
                source_incluse = True
            elif symbole == "H" and nombre_masse == 3:
                totals["tritium"][lieu]["total"] += masse_g
                totals["tritium"]["TOTAL"]["total"] += masse_g
                source_incluse = True
            elif symbole == "U" and nombre_masse == 233:
                totals["uranium_233"][lieu]["total"] += masse_g
                totals["uranium_233"]["TOTAL"]["total"] += masse_g
                source_incluse = True
            elif symbole == "U":
                masses_uranium[nombre_masse] = masses_uranium.get(nombre_masse, 0) + masse_g
                source_incluse = True

        # Uranium (hors 233) : à catégoriser globalement pour la source,
        # une fois tous ses isotopes d'uranium additionnés.
        if masses_uranium:
            masse_u_totale = sum(masses_uranium.values())
            masse_u235 = masses_uranium.get(235, 0.0)
            pct_u235 = (masse_u235 / masse_u_totale * 100) if masse_u_totale else 0.0
            categorie = categoriser_uranium(pct_u235)

            totals[categorie][lieu]["total"] += masse_u_totale
            totals[categorie]["TOTAL"]["total"] += masse_u_totale
            totals[categorie][lieu]["u235"] += masse_u235
            totals[categorie]["TOTAL"]["u235"] += masse_u235

        if source_incluse:
            sources_incluses += 1

    # Écriture dans le gabarit
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # Date de l'inventaire (celle de la génération du document, en face du
    # libellé "Inventaire du" en G3) -- demandé le 14/07/2026, absente
    # jusqu'ici.
    from datetime import date
    ws["H3"] = date.today().strftime("%d/%m/%Y")

    for categorie, (ligne, diviseur, fusionne) in LIGNES.items():
        for loc_nom, col_total in COLONNES_LOCALISATION.items():
            masse_totale = totals[categorie][loc_nom]["total"] / diviseur
            if masse_totale:
                ws[f"{col_total}{ligne}"] = round(masse_totale, 3)
            if not fusionne:  # lignes "Uranium enrichi" : 2e colonne = masse U235
                masse_235 = totals[categorie][loc_nom]["u235"] / diviseur
                col_u235 = chr(ord(col_total) + 1)  # C->D, E->F, G->H
                if masse_235:
                    ws[f"{col_u235}{ligne}"] = round(masse_235, 3)

    wb.save(output_path)
    return {"avertissements": avertissements, "sources_incluses": sources_incluses}
