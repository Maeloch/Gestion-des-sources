"""Tests de l'import des consommations historiques (transcrites depuis
des fiches manuscrites) -- demandé le 28/07/2026."""
import io
import openpyxl

from app.services.import_consommations_historiques import (
    importer_consommations_historiques,
    generer_modele_vide,
)
from app.services.units import quantite_restante_calculee
from app.repositories.source import SourceRepository


def _construire_fichier(lignes, entetes=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(entetes or [
        "ID Source", "Date", "Masse avant (g)", "Masse après (g)",
        "Quantité utilisée", "Commentaire", "Utilisateur",
    ])
    for ligne in lignes:
        ws.append(ligne)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _ecrire(tmp_path, lignes):
    chemin = tmp_path / "test.xlsx"
    chemin.write_bytes(_construire_fichier(lignes).read())
    return str(chemin)


def test_import_cree_les_consommations_avec_dates_historiques(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-HIST-TEST", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [["SRC-HIST-TEST", "10/03/2019", 200, 190, None, "premiere pesee", ""]])

    rapport = importer_consommations_historiques(chemin, db_session)

    assert rapport["consommations_creees"] == 1
    assert rapport["lignes_ignorees"] == []
    source = SourceRepository(db_session).get_by_id("SRC-HIST-TEST")
    conso = source.consumptions[0]
    assert (conso.timestamp.year, conso.timestamp.month, conso.timestamp.day) == (2019, 3, 10)
    assert conso.masse_avant == 200
    assert conso.masse_apres == 190


def test_ordre_des_lignes_dans_le_fichier_n_a_aucune_importance(admin_client, db_session, default_location, tmp_path):
    """Le point le plus important à vérifier : quantite_restante_calculee
    trie par date au moment du calcul, pas par ordre d'insertion -- donc
    un import dans le désordre chronologique (cas réaliste : les fiches
    papier ne sont pas forcément scannées/transcrites dans l'ordre) doit
    donner exactement le même résultat qu'un import trié."""
    admin_client.post("/sources/", json={
        "id": "SRC-HIST-DESORDRE", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [
        ["SRC-HIST-DESORDRE", "15/06/2020", 180, 175, None, "intermediaire", ""],
        ["SRC-HIST-DESORDRE", "10/03/2019", 200, 190, None, "premiere", ""],
        ["SRC-HIST-DESORDRE", "05/01/2021", 175, 170, None, "derniere", ""],
    ])

    rapport = importer_consommations_historiques(chemin, db_session)
    assert rapport["consommations_creees"] == 3

    source = SourceRepository(db_session).get_by_id("SRC-HIST-DESORDRE")
    resultat = quantite_restante_calculee(db_session, source)
    # Dernière pesée chronologique (05/01/2021) = 170 ; masse du récipient
    # déduite de la première pesée chronologique (10/03/2019) : 200-200=0.
    assert resultat == 170.0


def test_mode_quantite_directe_pour_source_gaz(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-HIST-GAZ-TEST", "type": "non-scellée", "etat_physique": "gaz",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 1000, "unite_quantite": "mL",
    })
    chemin = _ecrire(tmp_path, [["SRC-HIST-GAZ-TEST", "01/02/2020", None, None, 50, "purge", ""]])

    importer_consommations_historiques(chemin, db_session)
    source = SourceRepository(db_session).get_by_id("SRC-HIST-GAZ-TEST")
    assert quantite_restante_calculee(db_session, source) == 950.0


def test_masse_apres_absente_devient_pesee_de_controle(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-HIST-CTRL", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [["SRC-HIST-CTRL", "10/03/2019", 150, None, None, "pesee de controle", ""]])

    importer_consommations_historiques(chemin, db_session)
    source = SourceRepository(db_session).get_by_id("SRC-HIST-CTRL")
    conso = source.consumptions[0]
    assert conso.masse_apres == 150  # reprise telle quelle, rien prélevé
    assert conso.quantite_utilisee == 0


def test_ligne_avec_source_inexistante_ignoree_avec_raison(db_session, tmp_path):
    chemin = _ecrire(tmp_path, [["SRC-JAMAIS-CREEE", "10/03/2019", 200, 190, None, "", ""]])

    rapport = importer_consommations_historiques(chemin, db_session)
    assert rapport["consommations_creees"] == 0
    assert len(rapport["lignes_ignorees"]) == 1
    assert "introuvable" in rapport["lignes_ignorees"][0]["raison"]


def test_ligne_sans_date_ignoree_avec_raison(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-HIST-SANSDATE", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [["SRC-HIST-SANSDATE", None, 200, 190, None, "", ""]])

    rapport = importer_consommations_historiques(chemin, db_session)
    assert rapport["consommations_creees"] == 0
    assert "date" in rapport["lignes_ignorees"][0]["raison"]


def test_ligne_sans_aucune_quantite_ignoree_avec_raison(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-HIST-VIDE", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [["SRC-HIST-VIDE", "10/03/2019", None, None, None, "rien renseigne", ""]])

    rapport = importer_consommations_historiques(chemin, db_session)
    assert rapport["consommations_creees"] == 0
    assert "quantité" in rapport["lignes_ignorees"][0]["raison"]


def test_doublon_avec_existant_en_base_signale_mais_pas_bloque(admin_client, db_session, default_location, tmp_path):
    """31/07/2026, cas réel signalé : le même document importé deux fois
    a dupliqué tous les prélèvements. Doit être signalé clairement, mais
    ne doit PAS bloquer l'import (deux prélèvements réels le même jour
    restent possibles -- reste un import "idiot")."""
    admin_client.post("/sources/", json={
        "id": "SRC-DOUBLON-BASE", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    admin_client.post("/consumptions/", json={
        "source_id": "SRC-DOUBLON-BASE", "masse_avant": 200, "masse_apres": 190,
    })
    # Cette consommation existante est datée d'aujourd'hui (comportement
    # par défaut) -- le fichier réimporté vise donc la même date.
    from datetime import date
    chemin = _ecrire(tmp_path, [["SRC-DOUBLON-BASE", date.today().strftime("%d/%m/%Y"), 190, 180, None, "reimport par erreur", ""]])

    rapport = importer_consommations_historiques(chemin, db_session)

    assert rapport["consommations_creees"] == 1  # créée quand même, pas bloquée
    assert len(rapport["doublons_potentiels"]) == 1
    assert rapport["doublons_potentiels"][0]["source_id"] == "SRC-DOUBLON-BASE"


def test_doublon_entre_deux_lignes_du_meme_fichier_signale(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-DOUBLON-FICHIER", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [
        ["SRC-DOUBLON-FICHIER", "01/01/2021", 170, 165, None, "premiere ligne", ""],
        ["SRC-DOUBLON-FICHIER", "01/01/2021", 170, 165, None, "copier-coller en trop", ""],
    ])

    rapport = importer_consommations_historiques(chemin, db_session)

    assert rapport["consommations_creees"] == 2
    assert len(rapport["doublons_potentiels"]) == 1
    assert rapport["doublons_potentiels"][0]["ligne"] == 3  # la seconde ligne, pas la première


def test_meme_source_dates_differentes_pas_signalee_comme_doublon(admin_client, db_session, default_location, tmp_path):
    """Deux prélèvements réels à des dates différentes ne sont pas des
    doublons -- ne doivent jamais être signalés à tort."""
    admin_client.post("/sources/", json={
        "id": "SRC-PAS-DOUBLON", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [
        ["SRC-PAS-DOUBLON", "10/03/2019", 200, 190, None, "", ""],
        ["SRC-PAS-DOUBLON", "15/06/2020", 190, 180, None, "", ""],
    ])

    rapport = importer_consommations_historiques(chemin, db_session)

    assert rapport["consommations_creees"] == 2
    assert rapport["doublons_potentiels"] == []


def test_ligne_vide_ignoree_silencieusement(admin_client, db_session, default_location, tmp_path):
    """Une ligne complètement vide (ID Source absent) n'est pas une
    erreur -- probablement juste une ligne blanche dans le fichier."""
    admin_client.post("/sources/", json={
        "id": "SRC-HIST-BLANC", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [
        ["SRC-HIST-BLANC", "10/03/2019", 200, 190, None, "", ""],
        [None, None, None, None, None, None, None],
    ])

    rapport = importer_consommations_historiques(chemin, db_session)
    assert rapport["consommations_creees"] == 1
    assert rapport["lignes_ignorees"] == []


def test_genere_un_modele_vide_avec_les_bons_entetes(tmp_path):
    chemin = tmp_path / "modele.xlsx"
    generer_modele_vide(str(chemin))

    wb = openpyxl.load_workbook(str(chemin))
    ws = wb.active
    entetes = [c.value for c in ws[1]]
    assert entetes == [
        "ID Source", "Date", "Masse avant (g)", "Masse après (g)",
        "Quantité utilisée", "Commentaire", "Utilisateur",
    ]
    assert ws.max_row >= 2  # au moins une ligne d'exemple


def test_import_rattache_a_un_compte_existant_insensible_a_la_casse(admin_client, db_session, default_location, tmp_path):
    """31/08/2026 : GDO, BDL, CMO sont de vrais comptes connectés --
    l'import ne doit jamais leur créer un doublon historique, même si
    la fiche papier transcrite utilise une casse différente."""
    from app.repositories.user import UserRepository
    from app.models.user import UserCreate, UserRole
    bernadette = UserRepository(db_session).create(UserCreate(
        username="Bdl", email="bdl@test.fr", password="motdepasse123", role=UserRole.utilisateur,
    ))
    admin_client.post("/sources/", json={
        "id": "SRC-IMPORT-CASSE", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [["SRC-IMPORT-CASSE", "10/03/2019", 200, 190, None, "", "BDL"]])

    rapport = importer_consommations_historiques(chemin, db_session)

    assert rapport["utilisateurs_historiques_crees"] == []
    from app.models.consumption import ConsumptionDB
    conso = db_session.query(ConsumptionDB).filter(ConsumptionDB.source_id == "SRC-IMPORT-CASSE").first()
    assert conso.utilisateur_id == bernadette.id


def test_import_cree_un_enregistrement_historique_si_aucune_correspondance(admin_client, db_session, default_location, tmp_path):
    """Le cas réel signalé : Liatimi n'a pas de compte et n'en aura
    jamais, mais l'action qu'elle a réalisée doit rester rattachée."""
    admin_client.post("/sources/", json={
        "id": "SRC-IMPORT-LIATIMI", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [["SRC-IMPORT-LIATIMI", "10/03/2019", 200, 190, None, "", "Liatimi"]])

    rapport = importer_consommations_historiques(chemin, db_session)

    assert rapport["utilisateurs_historiques_crees"] == ["Liatimi"]
    from app.models.user import UserDB
    liatimi = db_session.query(UserDB).filter(UserDB.username == "Liatimi").first()
    assert liatimi is not None
    assert liatimi.is_active is False

    from app.models.consumption import ConsumptionDB
    conso = db_session.query(ConsumptionDB).filter(ConsumptionDB.source_id == "SRC-IMPORT-LIATIMI").first()
    assert conso.utilisateur_id == liatimi.id


def test_import_meme_utilisateur_sur_plusieurs_lignes_ne_cree_qu_un_seul_enregistrement(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-IMPORT-REPETE", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [
        ["SRC-IMPORT-REPETE", "10/03/2019", 200, 190, None, "", "Liatimi"],
        ["SRC-IMPORT-REPETE", "15/06/2019", 190, 180, None, "", "Liatimi"],
    ])

    importer_consommations_historiques(chemin, db_session)

    from app.models.user import UserDB
    assert db_session.query(UserDB).filter(UserDB.username == "Liatimi").count() == 1


def test_import_sans_utilisateur_renseigne_ne_cree_rien(admin_client, db_session, default_location, tmp_path):
    admin_client.post("/sources/", json={
        "id": "SRC-IMPORT-SANS-USER", "type": "non-scellée", "etat_physique": "liquide",
        "etat_utilisation": "en utilisation", "date_arrivee": "2018-01-01",
        "emplacement_habituel_id": default_location, "quantite_initiale": 200, "unite_quantite": "g",
    })
    chemin = _ecrire(tmp_path, [["SRC-IMPORT-SANS-USER", "10/03/2019", 200, 190, None, "", ""]])

    rapport = importer_consommations_historiques(chemin, db_session)

    assert rapport["consommations_creees"] == 1
    assert rapport["utilisateurs_historiques_crees"] == []
    from app.models.consumption import ConsumptionDB
    conso = db_session.query(ConsumptionDB).filter(ConsumptionDB.source_id == "SRC-IMPORT-SANS-USER").first()
    assert conso.utilisateur_id is None
    assert conso.utilisateur is None
